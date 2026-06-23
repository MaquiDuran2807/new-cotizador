import openpyxl
from openpyxl.styles import Font
from django.http import HttpResponse, HttpResponseRedirect
from django.shortcuts import render
from django.contrib import messages
from django.contrib.admin.helpers import ACTION_CHECKBOX_NAME
from django.db import transaction
from django.urls import path
from django.utils.html import format_html


class BulkActionsMixin:
    actions = ["assign_category_action", "upload_images_action"]

    def assign_category_action(self, request, queryset):
        selected = queryset.values_list("pk", flat=True)
        selected_str = ",".join(str(pk) for pk in selected)
        return HttpResponseRedirect(
            f"assign-category/?ids={selected_str}"
        )

    assign_category_action.short_description = "Asignar categoría a seleccionados"

    def upload_images_action(self, request, queryset):
        selected = queryset.values_list("pk", flat=True)
        selected_str = ",".join(str(pk) for pk in selected)
        return HttpResponseRedirect(
            f"upload-images/?ids={selected_str}"
        )

    upload_images_action.short_description = "Subir imágenes para seleccionados"

    def get_urls(self):
        urls = super().get_urls()
        info = self.model._meta.app_label, self.model._meta.model_name
        my_urls = [
            path(
                "assign-category/",
                self.admin_site.admin_view(self.assign_category_view),
                name="%s_%s_assign_category" % info,
            ),
            path(
                "upload-images/",
                self.admin_site.admin_view(self.upload_images_view),
                name="%s_%s_upload_images" % info,
            ),
        ]
        return my_urls + urls

    def assign_category_view(self, request):
        context = {
            **self.admin_site.each_context(request),
            "title": "Asignar categoría",
            "opts": self.model._meta,
            "media": self.media,
        }
        ids = request.GET.get("ids", "")
        pk_list = [pk for pk in ids.split(",") if pk]
        context["total"] = len(pk_list)

        if request.method == "POST":
            category_id = request.POST.get("category_id")
            if category_id and pk_list:
                from .models import Category
                try:
                    cat = Category.objects.get(pk=int(category_id))
                    updated = self.model.objects.filter(pk__in=pk_list).update(category=cat)
                    messages.success(request, f"{updated} productos actualizados a categoría '{cat.name}'")
                except (Category.DoesNotExist, ValueError):
                    messages.error(request, "Categoría inválida")
            return HttpResponseRedirect("../")

        categories = self._get_category_model().objects.all().values("id", "name")
        context["categories"] = categories
        return render(request, "admin/assign_category.html", context)

    def upload_images_view(self, request):
        context = {
            **self.admin_site.each_context(request),
            "title": "Subir imágenes para productos",
            "opts": self.model._meta,
            "media": self.media,
        }
        ids = request.GET.get("ids", "")
        pk_list = [pk for pk in ids.split(",") if pk]
        context["total"] = len(pk_list)

        if request.method == "POST":
            import zipfile, io, os
            from django.core.files.base import ContentFile

            zip_file = request.FILES.get("zip_file")
            if not zip_file:
                messages.error(request, "Debe seleccionar un archivo ZIP")
                return render(request, "admin/upload_images.html", context)

            import re
            success = 0
            errors = []
            zip_ids_found = set()
            try:
                with zipfile.ZipFile(zip_file) as zf:
                    for name in zf.namelist():
                        base = os.path.splitext(os.path.basename(name))[0]
                        if not base or not name.lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
                            continue
                        pk = None
                        # Try direct numeric ID
                        try:
                            pk = int(base)
                        except ValueError:
                            pass
                        # Try prod_XXXX_YY format
                        if pk is None:
                            m = re.match(r'prod_0*(\d+)_\d+', base, re.IGNORECASE)
                            if m:
                                pk = int(m.group(1))
                        if pk is None:
                            errors.append(f"'{base}' no es un ID numérico válido")
                            continue
                        zip_ids_found.add(pk)
                        if pk_list and pk not in [int(x) for x in pk_list]:
                            continue
                        try:
                            obj = self.model.objects.get(pk=pk)
                            img_data = zf.read(name)
                            ext = os.path.splitext(name)[1]
                            obj.image.save(f"{pk}{ext}", ContentFile(img_data), save=True)
                            success += 1
                        except self.model.DoesNotExist:
                            errors.append(f"Producto ID {pk} no encontrado")
                        except Exception as e:
                            errors.append(f"Error con ID {pk}: {e}")

                if success:
                    messages.success(request, f"{success} imágenes subidas correctamente")
                elif not errors:
                    sel_ids = sorted(int(x) for x in pk_list) if pk_list else []
                    zip_ids = sorted(zip_ids_found)
                    messages.info(request, f"IDs seleccionados: {sel_ids}. IDs encontrados en ZIP: {zip_ids}. "
                                         "Ninguno coincide. Revisá que los nombres de imagen correspondan a los IDs de tus productos.")
                if errors:
                    messages.warning(request, f"Ocurrieron {len(errors)} errores: {'; '.join(errors[:5])}")
            except zipfile.BadZipFile:
                messages.error(request, "Archivo ZIP inválido")

            return HttpResponseRedirect("../")

        return render(request, "admin/upload_images.html", context)

    def _get_category_model(self):
        from .models import Category
        return Category


class ThumbnailMixin:
    def show_image(self, obj):
        img_field = getattr(obj, 'image', None)
        if img_field and hasattr(img_field, 'url') and img_field.url:
            return format_html(
                '<img src="{}" style="width:50px;height:50px;object-fit:cover;border-radius:4px;border:1px solid #ddd;" />',
                img_field.url
            )
        return "—"
    show_image.short_description = "Imagen"


class ExportExcelMixin(ThumbnailMixin, BulkActionsMixin):
    actions = ["download_excel_action", "assign_category_action", "upload_images_action"]
    change_list_template = "admin/change_list_with_import.html"

    def get_urls(self):
        urls = super().get_urls()
        info = self.model._meta.app_label, self.model._meta.model_name
        my_urls = [
            path(
                "import-excel/",
                self.admin_site.admin_view(self.import_excel_view),
                name="%s_%s_import" % info,
            ),
        ]
        return my_urls + urls

    def get_export_filename(self, request):
        return f"{self.model._meta.verbose_name_plural}.xlsx"

    def _get_field_label(self, field_name):
        if field_name == "__str__":
            return str(self.model._meta.verbose_name)
        if hasattr(self, field_name):
            method = getattr(self, field_name)
            if hasattr(method, "short_description"):
                return method.short_description
        try:
            field = self.model._meta.get_field(field_name)
            return field.verbose_name
        except:
            pass
        return field_name

    def _get_field_value(self, obj, field_name):
        if field_name == "__str__":
            return str(obj)
        if hasattr(self, field_name):
            method = getattr(self, field_name)
            if callable(method):
                return str(method(obj))
        if "__" in field_name:
            parts = field_name.split("__")
            value = obj
            for part in parts:
                try:
                    value = getattr(value, part)
                except AttributeError:
                    return ""
            return str(value) if value is not None else ""
        try:
            value = getattr(obj, field_name)
            if callable(value):
                value = value()
            return str(value) if value is not None else ""
        except:
            return ""

    def download_excel_action(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        model_name = str(self.model._meta.verbose_name_plural)
        ws.title = model_name[:31]

        fields = self.list_display
        header_font = Font(bold=True, size=11)

        for col, field_name in enumerate(fields, 1):
            cell = ws.cell(row=1, column=col, value=self._get_field_label(field_name))
            cell.font = header_font

        for row_num, obj in enumerate(queryset, 2):
            for col_num, field_name in enumerate(fields, 1):
                ws.cell(row=row_num, column=col_num, value=self._get_field_value(obj, field_name))

        for col_cells in ws.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 3, 60)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{self.get_export_filename(request)}"'
        wb.save(response)
        return response

    download_excel_action.short_description = "Descargar seleccionados como Excel"

    def import_excel_view(self, request):
        context = {
            **self.admin_site.each_context(request),
            "title": f"Importar {self.model._meta.verbose_name_plural}",
            "opts": self.model._meta,
            "media": self.media,
            "has_add_permission": self.has_add_permission(request),
            "has_change_permission": self.has_change_permission(request),
            "has_delete_permission": self.has_delete_permission(request),
            "has_view_permission": self.has_view_permission(request),
        }

        if request.method == "POST":
            excel_file = request.FILES.get("excel_file")
            if not excel_file:
                messages.error(request, "Debe seleccionar un archivo Excel.")
                return render(request, "admin/import_excel.html", context)

            if not excel_file.name.endswith((".xlsx", ".xls")):
                messages.error(request, "Solo se permiten archivos .xlsx o .xls.")
                return render(request, "admin/import_excel.html", context)

            try:
                wb = openpyxl.load_workbook(excel_file)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))

                if len(rows) < 2:
                    messages.error(request, "El archivo debe tener encabezados y al menos una fila de datos.")
                    return render(request, "admin/import_excel.html", context)

                headers = [str(h).strip() if h is not None else "" for h in rows[0]]
                data_rows = rows[1:]

                created, updated, errors = self._process_import_rows(headers, data_rows)

                if created:
                    messages.success(request, f"{created} registros creados con éxito.")
                if updated:
                    messages.success(request, f"{updated} registros actualizados con éxito.")
                if errors:
                    messages.warning(request, f"Se completó con {len(errors)} errores.")
                    context["import_errors"] = errors

                if not created and not updated and not errors:
                    messages.info(request, "No se procesaron filas.")

            except Exception as e:
                messages.error(request, f"Error al procesar el archivo: {str(e)}")

        return render(request, "admin/import_excel.html", context)

    def _process_import_rows(self, headers, data_rows):
        created = 0
        updated = 0
        errors = []

        model_fields = {f.name: f for f in self.model._meta.fields}

        col_map = {}
        for col_idx, header in enumerate(headers):
            hl = header.lower()
            match = None

            # 1) Match exact field name
            for fname, field in model_fields.items():
                if hl == fname.lower():
                    match = fname
                    break

            # 2) Match by verbose_name
            if not match:
                for fname, field in model_fields.items():
                    vn = str(field.verbose_name).lower() if field.verbose_name else ""
                    if hl == vn:
                        match = fname
                        break

            # 3) Match "category_id" → FK field "category"
            if not match and hl.endswith("_id"):
                base = hl[:-3]
                for fname, field in model_fields.items():
                    if field.is_relation and fname.lower() == base:
                        match = fname
                        break

            # 4) Explicit "id"
            if not match and hl == "id":
                match = "id"

            col_map[col_idx] = match

        for row_idx, row in enumerate(data_rows, 2):
            try:
                obj_data = {}
                skip_row = False
                for col_idx, value in enumerate(row):
                    field_name = col_map.get(col_idx)
                    if not field_name:
                        continue
                    if value is None or str(value).strip() == "":
                        continue

                    field = model_fields.get(field_name)
                    if not field:
                        obj_data[field_name] = str(value)
                        continue

                    ftype = field.get_internal_type()
                    val = self._convert_value(value, ftype, field, row_idx, errors)
                    if val is None:
                        skip_row = True
                        break
                    obj_data[field_name] = val

                if skip_row or not obj_data:
                    continue

                with transaction.atomic():
                    obj_id = obj_data.pop("id", None)
                    if obj_id:
                        try:
                            instance = self.model.objects.get(pk=obj_id)
                            for key, val in obj_data.items():
                                setattr(instance, key, val)
                            instance.save()
                            updated += 1
                            continue
                        except self.model.DoesNotExist:
                            obj_data["id"] = obj_id
                    self.model.objects.create(**obj_data)
                    created += 1

            except Exception as e:
                errors.append(f"Fila {row_idx}: {str(e)}")

        return created, updated, errors

    def _convert_value(self, value, ftype, field, row_idx, errors):
        if ftype == "ForeignKey":
            resolved = self._resolve_fk(value, field)
            if resolved is None:
                errors.append(f"Fila {row_idx}: No se encontró '{value}' para '{field.name}'")
            return resolved
        if ftype in ("IntegerField", "AutoField", "BigIntegerField", "SmallIntegerField", "PositiveIntegerField", "PositiveSmallIntegerField", "BigAutoField", "SmallAutoField"):
            try:
                return int(float(str(value).replace(",", ".")))
            except (ValueError, TypeError):
                errors.append(f"Fila {row_idx}: '{value}' no es un número válido para '{field.name}'")
                return None
        if ftype == "FloatField":
            try:
                return float(str(value).replace(",", "."))
            except (ValueError, TypeError):
                errors.append(f"Fila {row_idx}: '{value}' no es un número decimal válido para '{field.name}'")
                return None
        if ftype in ("BooleanField", "NullBooleanField"):
            if isinstance(value, bool):
                return value
            return str(value).lower() in ("si", "yes", "true", "1", "verdadero")
        if ftype in ("CharField", "TextField", "SlugField", "URLField", "EmailField"):
            return str(value)
        if ftype == "ImageField":
            return None
        return str(value)

    def _resolve_fk(self, value, field):
        related = field.remote_field.model
        # Try numeric ID first
        try:
            return related.objects.get(pk=int(float(str(value))))
        except (ValueError, TypeError, OverflowError):
            pass
        except related.DoesNotExist:
            # ID exists but not found — fall through to name match
            pass
        # Try by common string fields
        for str_field in ("name", "title", "code", "nombre"):
            if hasattr(related, str_field):
                kwargs = {f"{str_field}__iexact": str(value).strip()}
                obj = related.objects.filter(**kwargs).first()
                if obj:
                    return obj
        # Try partial match on name
        for str_field in ("name", "title", "nombre"):
            if hasattr(related, str_field):
                kwargs = {f"{str_field}__icontains": str(value).strip()}
                obj = related.objects.filter(**kwargs).first()
                if obj:
                    return obj
        return None
